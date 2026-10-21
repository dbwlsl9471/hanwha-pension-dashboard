# -*- coding: utf-8 -*-
"""포털 원본 엑셀 → js/data-store.js 변환"""
import json
import re
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

BASE = Path(__file__).resolve().parent
OUT = BASE / "js" / "data-store.js"

SKIP_TYPES = {"합계", "자사계열사", "기타 사업자"}
TYPE_MAP = {
    "예금성 원리금보장": "예금성",
    "시장성 원리금보장": "시장성",
    "원리금비보장": "실적배당형",
    "원리금 비보장": "실적배당형",
    # 예·시장 분리 이전 공시
    "원리금 보장": "원리금보장",
    "원리금보장": "원리금보장",
}


def sector_of(name: str) -> str:
    n = name.replace(" ", "")
    if "은행" in n or n.endswith("뱅크") or "iM뱅크" in n:
        return "은행"
    if "증권" in n:
        return "증권"
    if "생명" in n or "라이프" in n or "연금보험" in n:
        return "생보"
    if "손해" in n or "화재" in n or "해상" in n:
        return "손보"
    return "기타"


def to_num(v):
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("%", "").strip()
    if s in ("", "-", "—"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_quarter_code(stem: str):
    """수익률244Q → 2024Q4"""
    m = re.search(r"(\d{2})(\d)Q", stem)
    if not m:
        return None
    yy, q = int(m.group(1)), int(m.group(2))
    year = 2000 + yy if yy < 100 else yy
    return f"{year}Q{q}"


def parse_returns(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    # year/quarter from header or filename
    year = ws.cell(3, 2).value
    q = ws.cell(3, 4).value
    if year and q:
        key = f"{int(year)}Q{int(q)}"
    else:
        key = parse_quarter_code(path.stem)
    rows = []
    cur = None
    for r in range(10, ws.max_row + 1):
        name = ws.cell(r, 1).value
        typ_raw = ws.cell(r, 2).value
        if name:
            cur = str(name).strip()
        if not cur or not typ_raw:
            continue
        typ_raw = str(typ_raw).strip()
        if typ_raw in SKIP_TYPES:
            continue
        typ = TYPE_MAP.get(typ_raw)
        if not typ:
            continue

        def block(c0):
            return {
                "aum": to_num(ws.cell(r, c0).value),
                "r1": to_num(ws.cell(r, c0 + 1).value),
                "r3": to_num(ws.cell(r, c0 + 2).value),
                "r5": to_num(ws.cell(r, c0 + 3).value),
                "r7": to_num(ws.cell(r, c0 + 4).value),
                "r10": to_num(ws.cell(r, c0 + 5).value),
            }

        rows.append(
            {
                "name": cur,
                "sector": sector_of(cur),
                "type": typ,
                "DB": block(3),
                "DC": block(9),
                "IRP": block(15),
            }
        )
    return key, rows


def parse_default(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    key = parse_quarter_code(path.stem)
    # try header
    for c in range(1, 10):
        v = ws.cell(3, c).value
        if v and "분기" in str(v):
            pass
    rows = []
    for r in range(9, ws.max_row + 1):
        name = ws.cell(r, 2).value
        risk = ws.cell(r, 5).value
        if not name or not risk:
            continue
        name = str(name).strip()
        risk = str(risk).strip()
        rows.append(
            {
                "name": name,
                "sector": sector_of(name),
                "product": str(ws.cell(r, 1).value or "").strip(),
                "risk": risk,
                "aum_dc": to_num(ws.cell(r, 6).value),
                "aum_irp": to_num(ws.cell(r, 7).value),
                "aum": to_num(ws.cell(r, 8).value),
                "r1m": to_num(ws.cell(r, 9).value),
                "r3m": to_num(ws.cell(r, 10).value),
                "r6m": to_num(ws.cell(r, 11).value),
                "r1y": to_num(ws.cell(r, 12).value),
                "r3y": to_num(ws.cell(r, 13).value),
            }
        )
    return key, rows


def parse_fee(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    year = ws.cell(3, 2).value
    key = str(int(year)) if year else re.search(r"(\d{4})", path.stem).group(1)
    rows = []
    for r in range(10, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        name = str(name).strip()

        def plan(c0):
            # 합계 대면/비대면 at c0+4, c0+5 (0-based offset from plan start)
            # plan starts: DB=2, DC=8, IRP=14 → face total at +4, online at +5
            face = to_num(ws.cell(r, c0 + 4).value)
            online = to_num(ws.cell(r, c0 + 5).value)
            # prefer 대면, else 비대면
            rate = face if face is not None else online
            return {
                "face": face,
                "online": online,
                "rate": rate,
                "metric": "수수료 비용부담률",
            }

        rows.append(
            {
                "name": name,
                "sector": sector_of(name),
                "DB": plan(2),
                "DC": plan(8),
                "IRP": plan(14),
            }
        )
    return key, rows


def main():
    store = {
        "meta": {
            "company": "한화투자증권",
            "competitors_default": [
                "미래에셋증권",
                "삼성증권",
                "한국투자증권",
                "NH투자증권",
                "현대차증권",
            ],
            "updatedAt": None,
        },
        "returns": {},
        "defaults": {},
        "fees": {},
        "sources": {},
    }

    for path in sorted(BASE.glob("수익률*.xlsx")):
        key, rows = parse_returns(path)
        store["returns"][key] = rows
        store["sources"][f"returns:{key}"] = {
            "file": path.name,
            "kind": "공시",
            "rows": len(rows),
        }
        print(f"returns {key}: {len(rows)} rows from {path.name}")

    for path in sorted(BASE.glob("디폴트옵션*.xlsx")):
        key, rows = parse_default(path)
        if not key:
            print("skip default", path.name)
            continue
        store["defaults"][key] = rows
        store["sources"][f"defaults:{key}"] = {
            "file": path.name,
            "kind": "공시",
            "rows": len(rows),
        }
        print(f"defaults {key}: {len(rows)} rows from {path.name}")

    for path in sorted(BASE.glob("수수료*.xlsx")):
        key, rows = parse_fee(path)
        store["fees"][key] = rows
        store["sources"][f"fees:{key}"] = {
            "file": path.name,
            "kind": "공시",
            "rows": len(rows),
        }
        print(f"fees {key}: {len(rows)} rows from {path.name}")

    from datetime import datetime

    store["meta"]["updatedAt"] = datetime.now().isoformat(timespec="seconds")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(store, ensure_ascii=False)
    OUT.write_text(
        "/* auto-generated by build_data.py — do not edit by hand */\n"
        f"window.APP_DATA = {payload};\n",
        encoding="utf-8",
    )
    print("wrote", OUT, "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
